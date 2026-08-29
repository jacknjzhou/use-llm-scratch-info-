"""文件解析与分块测试：txt/docx/xlsx/pdf(文本层)/扫描件OCR兜底/分块逻辑。"""
import pytest

from app.parsers import parse_file


def _chunked_text(paragraphs=1, chars=200):
    return "\n".join(("内容" * (chars // 2 + 1))[:chars] for _ in range(paragraphs))


def test_txt_parse_and_chunk(tmp_path):
    p = tmp_path / "sample.txt"
    text = _chunked_text(paragraphs=10, chars=2000)
    p.write_text(text, encoding="utf-8")
    chunks = parse_file(str(p))
    assert len(chunks) >= 2                       # 长文本被分块
    assert chunks[0]["index"] == 0
    assert chunks[0]["page"] == 1
    assert all(c["text"] for c in chunks)          # 无空块
    # 分块连续性：第二块应包含第一块尾部重叠内容
    assert chunks[1]["text"].startswith(chunks[0]["text"][-200:])


def test_docx_parse(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("甲方：测试科技有限公司")
    d.add_paragraph("合同金额：人民币120万元")
    path = tmp_path / "c.docx"
    d.save(str(path))
    chunks = parse_file(str(path))
    assert any("甲方：测试科技有限公司" in c["text"] for c in chunks)


def test_xlsx_parse(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "报销单"
    ws.append(["姓名", "金额"])
    ws.append(["张三", "553.0"])
    path = tmp_path / "r.xlsx"
    wb.save(str(path))
    chunks = parse_file(str(path))
    assert any("张三" in c["text"] and "报销单" in c["text"] for c in chunks)


def test_pdf_with_text_layer(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "高铁票 G101 北京南 到 上海虹桥 二等座 553元")
    path = tmp_path / "ticket.pdf"
    doc.save(str(path))
    doc.close()
    chunks = parse_file(str(path))
    assert any("G101" in c["text"] for c in chunks)


def test_scan_pdf_ocr_fallback(tmp_path):
    """无文本层 PDF：OCR 可用则返回文本，否则必须报带指引的错误。"""
    import pymupdf
    import pytesseract

    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect((50, 50, 200, 200))
    path = tmp_path / "scan.pdf"
    doc.save(str(path))
    doc.close()

    try:
        pytesseract.get_tesseract_version()
    except pytesseract.TesseractNotFoundError:
        with pytest.raises(ValueError, match="OCR"):
            parse_file(str(path))
    else:
        chunks = parse_file(str(path))
        assert isinstance(chunks, list)


def test_unsupported_extension(tmp_path):
    p = tmp_path / "x.exe"
    p.write_bytes(b"xx")
    with pytest.raises(ValueError, match="暂不支持"):
        parse_file(str(p))
